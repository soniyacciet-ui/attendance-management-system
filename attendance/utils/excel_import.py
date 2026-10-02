"""
Excel bulk import utilities for the Attendance Management System.

Design rules:
  * No database writes here.
  * No HTTP handling here.
  * Views call parse_and_validate(), get back structured rows,
    then commit inside transaction.atomic().
"""

from dataclasses import dataclass, field
from datetime import datetime, date
from io import BytesIO
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ======================================================================
# CONSTANTS
# ======================================================================

REQUIRED_HEADERS = [
    "Student ID",
    "Student Name",
    "Email",
    "Phone",
    "Gender",
    "DOB",
    "Department",
    "Year",
    "Section",
    "Parent 1 Name",
    "Parent 1 Phone",
    "Parent 1 Email",
    "Parent 1 Relation",
    "Parent 2 Name",
    "Parent 2 Phone",
    "Parent 2 Email",
    "Parent 2 Relation",
]

VALID_GENDERS = {"MALE", "FEMALE", "OTHER"}
VALID_RELATIONS = {"FATHER", "MOTHER", "GUARDIAN"}
MAX_PARENTS = 2


# ======================================================================
# DATA CLASSES
# ======================================================================

@dataclass
class ParentRow:
    name: str = ""
    phone: str = ""
    email: str = ""
    relationship: str = "FATHER"


@dataclass
class StudentRow:
    row_number: int
    register_number: str = ""
    name: str = ""
    email: str = ""
    phone: str = ""
    gender: str = ""
    dob: Optional[date] = None
    department_name: str = ""
    year: Optional[int] = None
    section: str = ""
    parents: list = field(default_factory=list)   # list[ParentRow]
    errors: list = field(default_factory=list)     # list[str]

    @property
    def is_valid(self):
        return not self.errors


# ======================================================================
# LOW-LEVEL HELPERS
# ======================================================================

def _s(value) -> str:
    """Convert any cell value to a clean string."""
    if value is None:
        return ""
    return str(value).strip()


def _parse_date(value) -> Optional[date]:
    """Accept datetime, date, or string in several common formats."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognised date format: {s!r}")


def _normalise_relation(value: str) -> str:
    v = value.strip().upper()
    aliases = {
        "FATHER": "FATHER", "DAD": "FATHER", "F": "FATHER",
        "MOTHER": "MOTHER", "MOM": "MOTHER", "M": "MOTHER",
        "GUARDIAN": "GUARDIAN", "G": "GUARDIAN", "OTHER": "GUARDIAN",
    }
    return aliases.get(v, v)


def _is_valid_email(email: str) -> bool:
    if not email:
        return True
    return "@" in email and "." in email.split("@")[-1] and " " not in email


def _is_valid_phone(phone: str) -> bool:
    if not phone:
        return True
    digits = "".join(c for c in phone if c.isdigit())
    return 7 <= len(digits) <= 15


# ======================================================================
# TEMPLATE BUILDER
# ======================================================================

def build_template_workbook() -> BytesIO:
    """
    Return a BytesIO containing the ready-to-fill .xlsx template.
      Sheet 1: Students   (headers + one example row)
      Sheet 2: Instructions
    """
    wb = Workbook()

    # ----- Sheet 1 -----
    ws = wb.active
    ws.title = "Students"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0B2B4A")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, header in enumerate(REQUIRED_HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center

    # One clearly-labelled example row
    example = [
        "STU001", "Arun Kumar", "arun@example.com", "9876543210",
        "MALE", "2004-05-15", "AI&DS", 3, "A",
        "Ravi Kumar", "9876543211", "ravi@example.com", "FATHER",
        "Priya Kumar", "9876543212", "priya@example.com", "MOTHER",
    ]
    for col, value in enumerate(example, start=1):
        ws.cell(row=2, column=col, value=value)

    # Column widths
    widths = [12, 22, 26, 14, 10, 12, 18, 8, 10,
              20, 14, 26, 14, 20, 14, 26, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    # ----- Sheet 2 -----
    ins = wb.create_sheet("Instructions")
    lines = [
        ("Bulk Student Import — Instructions", True),
        ("", False),
        ("1. Fill in the 'Students' sheet. Do not rename or reorder columns.", False),
        ("2. Delete the example row (row 2) before uploading, or overwrite it.", False),
        ("3. Student ID and Student Name are required.", False),
        ("4. Department must already exist in your college.", False),
        ("5. A student can have at most 2 parents.", False),
        ("6. Parent Email is required whenever any parent field is filled.", False),
        ("7. Relation must be one of: FATHER, MOTHER, GUARDIAN.", False),
        ("8. Gender must be one of: MALE, FEMALE, OTHER (optional).", False),
        ("9. DOB format: YYYY-MM-DD (optional).", False),
        ("10. Year must be a number 1–6 (optional).", False),
        ("11. Save as .xlsx and upload from the Bulk Add Students page.", False),
    ]
    for i, (text, is_title) in enumerate(lines, start=1):
        c = ins.cell(row=i, column=1, value=text)
        if is_title:
            c.font = Font(bold=True, size=14, color="0B2B4A")
    ins.column_dimensions["A"].width = 100

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ======================================================================
# PARSER + VALIDATOR
# ======================================================================

def parse_and_validate(
    file_bytes: bytes,
    *,
    existing_register_numbers: set,
    department_lookup: dict,   # {lowercased_name: Department}
) -> tuple[list[StudentRow], list[str]]:
    """
    Parse the uploaded .xlsx and validate every row.

    Returns
    -------
    rows          : list[StudentRow]  — each may carry .errors
    global_errors : list[str]         — file-level problems (bad headers, etc.)
    """
    global_errors: list[str] = []

    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as e:
        return [], [f"Could not read the Excel file: {e}"]

    if "Students" in wb.sheetnames:
        ws = wb["Students"]
    else:
        ws = wb[wb.sheetnames[0]]

    if ws.max_row < 2:
        return [], ["The uploaded file contains no data rows."]

    # ---------- Header check ----------
    header_row = [_s(ws.cell(row=1, column=c).value)
                  for c in range(1, ws.max_column + 1)]
    missing = [h for h in REQUIRED_HEADERS if h not in header_row]
    if missing:
        return [], ["Missing required columns: " + ", ".join(missing)]

    col = {h: header_row.index(h) + 1 for h in REQUIRED_HEADERS}

    rows: list[StudentRow] = []
    seen_in_file: dict[str, int] = {}

    for r in range(2, ws.max_row + 1):

        def get(header):
            return _s(ws.cell(row=r, column=col[header]).value)

        # Skip fully empty rows
        if not any(get(h) for h in REQUIRED_HEADERS):
            continue

        row = StudentRow(row_number=r)
        row.register_number = get("Student ID")
        row.name = get("Student Name")
        row.email = get("Email")
        row.phone = get("Phone")
        row.gender = get("Gender").upper()
        row.department_name = get("Department")
        row.section = get("Section")

        # ---- Student ID ----
        if not row.register_number:
            row.errors.append("Student ID is missing.")

        # ---- Student Name ----
        if not row.name:
            row.errors.append("Student Name is missing.")

        # ---- Duplicate inside the file ----
        if row.register_number:
            key = row.register_number.lower()
            if key in seen_in_file:
                row.errors.append(
                    f"Duplicate Student ID {row.register_number} "
                    f"(already used on row {seen_in_file[key]})."
                )
            else:
                seen_in_file[key] = r

        # ---- Duplicate in database ----
        if row.register_number and row.register_number in existing_register_numbers:
            row.errors.append(
                f"Student ID {row.register_number} already exists in the database."
            )

        # ---- Email ----
        if row.email and not _is_valid_email(row.email):
            row.errors.append(f"Invalid student email: {row.email}")

        # ---- Phone ----
        if row.phone and not _is_valid_phone(row.phone):
            row.errors.append(f"Invalid student phone: {row.phone}")

        # ---- Gender ----
        if row.gender and row.gender not in VALID_GENDERS:
            row.errors.append(
                f"Invalid Gender '{row.gender}' "
                f"(use MALE / FEMALE / OTHER)."
            )

        # ---- DOB ----
        dob_raw = get("DOB")
        if dob_raw:
            try:
                row.dob = _parse_date(dob_raw)
            except ValueError as e:
                row.errors.append(str(e))

        # ---- Year ----
        year_raw = get("Year")
        if year_raw:
            try:
                row.year = int(float(year_raw))
                if row.year < 1 or row.year > 6:
                    row.errors.append(
                        f"Year '{year_raw}' is out of range (1–6)."
                    )
            except ValueError:
                row.errors.append(f"Year '{year_raw}' is not a valid number.")

        # ---- Department ----
        if row.department_name:
            if row.department_name.strip().lower() not in department_lookup:
                row.errors.append(
                    f"Department '{row.department_name}' not found in this college."
                )
        else:
            row.errors.append("Department is missing.")

        # ---- Parents ----
        parents: list[ParentRow] = []
        for i in (1, 2):
            p_name = get(f"Parent {i} Name")
            p_phone = get(f"Parent {i} Phone")
            p_email = get(f"Parent {i} Email")
            p_rel = get(f"Parent {i} Relation")

            if not any([p_name, p_phone, p_email, p_rel]):
                continue

            p = ParentRow(
                name=p_name,
                phone=p_phone,
                email=p_email,
                relationship=_normalise_relation(p_rel) if p_rel else "FATHER",
            )

            if not p.name:
                row.errors.append(
                    f"Parent {i} Name is empty but other Parent {i} fields are filled."
                )
            if not p.email:
                row.errors.append(
                    f"Parent {i} Email is required (used for notifications)."
                )
            elif not _is_valid_email(p.email):
                row.errors.append(f"Invalid Parent {i} email: {p.email}")
            if p.phone and not _is_valid_phone(p.phone):
                row.errors.append(f"Invalid Parent {i} phone: {p.phone}")
            if p.relationship not in VALID_RELATIONS:
                row.errors.append(
                    f"Parent {i} Relation '{p_rel}' is invalid "
                    f"(use FATHER / MOTHER / GUARDIAN)."
                )

            parents.append(p)

        # Defensive cap (two columns already enforce this)
        if len(parents) > MAX_PARENTS:
            row.errors.append(f"More than {MAX_PARENTS} parents provided.")

        row.parents = parents[:MAX_PARENTS]
        rows.append(row)

    return rows, global_errors


# ======================================================================
# ERROR-REPORT BUILDER
# ======================================================================

def build_error_report(rows: list[StudentRow]) -> BytesIO:
    """
    Return a BytesIO with an .xlsx containing only the invalid rows.
    Columns: Original Row | Student ID | Errors
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Errors"

    headers = ["Original Row", "Student ID", "Errors"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="B02A37")

    r = 2
    for row in rows:
        if row.is_valid:
            continue
        ws.cell(row=r, column=1, value=row.row_number)
        ws.cell(row=r, column=2, value=row.register_number)
        ws.cell(row=r, column=3, value=" | ".join(row.errors))
        r += 1

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 100

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf